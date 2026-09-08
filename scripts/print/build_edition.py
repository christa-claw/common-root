#!/usr/bin/env python3
"""
build_edition.py — pour a whole edition into a printer's own template.

WHY THIS EXISTS
---------------
build_print_sample.py answers "what is this?" in four pages, set in our
typography. This answers the question that follows a yes: "send us your
template and we will return your book, in your fonts, with your chapter
separators, in whichever ordering you are selling."

The shop's design stays theirs. We never invent a font, a rule, an ornament or
a margin — we emit text tagged with THEIR style names and let their template
resolve it. That is the whole architectural idea, and it is why there is a
style map rather than a stylesheet.

FORMATS
-------
tagged  InDesign Tagged Text. What a book typesetter will actually ask for.
        Place it into their existing InDesign document and every paragraph
        picks up the shop's paragraph and character styles — fonts, chapter
        separators, drop caps, the lot. Emitted UTF-16LE as InDesign requires.
html    Their HTML/CSS shell with @font-face and ornament markup. For shops
        that do not typeset in-house, and for anything that must print from a
        browser. Fragment-based, so their separator markup is reproduced
        verbatim once per chapter.
docx    Their .docx, opened as a template; paragraphs are appended carrying
        their named styles. Common at small print-on-demand outfits.
idml    Not implemented on purpose — see the IdmlWriter docstring.

SCOPE
-----
Whole Bible by default, because that is what a Bible printer prints. --books
exists for a taster in their fonts to attach to a reply, not as the main use.

Usage:
    python3 scripts/print/build_edition.py --self-test
    python3 scripts/print/build_edition.py --translation bible-kjv-1611 \\
        --ordering chronological --format tagged \\
        --styles templates/acme-styles.yml --out out/acme-kjv-chrono.txt
    python3 scripts/print/build_edition.py --translation bible-kjv-1611 \\
        --ordering canonical --format html \\
        --template templates/example.html --out out/kjv.html --books GEN,EXO
"""

import argparse
import base64
import html as htmlmod
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request

BASEX = os.environ.get("BASEX_URL", "http://localhost:8984/rest/religioustext")
AUTH = base64.b64encode(
    os.environ.get("BASEX_AUTH", "admin:admin").encode()).decode()
NS = "declare namespace rt='http://religioustext.org/schema/1.0'; "

# The corpus stamps one global sequence attribute per ordering. They are
# denormalised onto the verse precisely so a whole edition can be streamed in
# any order without walking the tree.
ORDERINGS = {
    "canonical":     "globalCanonicalSeq",
    "chronological": "globalChronologicalSeq",
    "narrative":     "globalNarrativeSeq",
    "tanakh":        "globalTanakhSeq",
}

# Every role a template may style. A shop that has no ornament simply omits
# `separator` and nothing is emitted for it.
ROLES = ("book_title", "chapter_head", "separator", "body", "verse_number")

CHUNK = 2000            # verses per BaseX request


# ── corpus ────────────────────────────────────────────────────────────────────

class CorpusError(RuntimeError):
    pass


def xq(query, timeout=120):
    url = f"{BASEX}?{urllib.parse.urlencode({'query': query})}"
    req = urllib.request.Request(url, headers={"Authorization": f"Basic {AUTH}",
                                               "Accept": "text/plain"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.read().decode("utf-8")
    except urllib.error.URLError as e:
        raise CorpusError(f"BaseX at {BASEX} did not answer: {e}") from e


def available_orderings(doc):
    """Which sequence attributes are actually stamped on this document.

    Asked before anything else because an unstamped ordering is the failure
    that has bitten this project twice: the stamper reports success, stamps
    zero verses, and the defect only surfaces much later as text in the wrong
    order. Here it surfaces immediately, by name.
    """
    found = {}
    for name, attr in ORDERINGS.items():
        q = (NS + f"count(db:open('religioustext', '{doc}.xml')"
                  f"//rt:verse[@{attr}])")
        # A CorpusError here is BaseX being unreachable, NOT an unstamped
        # ordering, and must not be swallowed: reporting "no verses stamped"
        # when the server is simply down sends the reader to the stamper for a
        # problem that is a closed port. Only an unparseable count is a zero.
        try:
            found[name] = int(xq(q, timeout=60).strip() or 0)
        except ValueError:
            found[name] = 0
    return found


def stream_verses(doc, ordering, books=None, limit=None):
    """Yield (book_name, chapter, verse_number, text) in the chosen order."""
    attr = ORDERINGS[ordering]
    counts = available_orderings(doc)
    total = counts.get(ordering, 0)
    if total == 0:
        stamped = ", ".join(f"{k} ({v})" for k, v in counts.items() if v) or "none"
        raise CorpusError(
            f"'{doc}' has no verses stamped with @{attr}, so the {ordering} "
            f"ordering cannot be produced. Stamped orderings: {stamped}. "
            f"Run the matching stamper (e.g. stamp_chronological.py "
            f"--translation {doc}) and check it reports a non-zero count.")

    book_pred = ""
    if books:
        wanted = " or ".join(f"$v/@bookName = '{b}'" for b in books)
        book_pred = f" and ({wanted})"

    seen = 0
    start = 1
    while True:
        end = start + CHUNK - 1
        q = (NS +
             f"for $v in db:open('religioustext', '{doc}.xml')//rt:verse[@{attr}] "
             f"let $s := xs:integer($v/@{attr}) "
             f"where $s >= {start} and $s <= {end}{book_pred} "
             f"order by $s "
             f"return string-join((string($v/@bookName), "
             f"string($v/@chapterNumber), string($v/@number), "
             f"normalize-space(string($v))), '&#9;')")
        block = xq(q)
        rows = [ln for ln in block.split("\n") if ln.count("\t") >= 3]
        for ln in rows:
            book, chap, num, text = ln.split("\t", 3)
            yield book, chap, num, clean(text)
            seen += 1
            if limit and seen >= limit:
                return
        start = end + 1
        if start > total * 2:      # every sequence exhausted, with slack for gaps
            return
        if seen >= total:
            return


# Leaked translator footnotes: the wldeh source JSONs embed notes inline, keyed
# "<chapter>.<verse> <lemma>: <note>", and the ingest kept them. See the INLINE
# FOOTNOTE LEAK backlog note (2026-07-24). This strips them on the way out so a
# printer is never sent broken text. It is a WORKAROUND: the corpus, the reader
# and the public XML download all still carry the corruption.
FOOTNOTE_TAIL = re.compile(r"\s*\b\d+\.\d+\s+\S[^:]{0,40}:.*$")


def clean(text):
    return FOOTNOTE_TAIL.sub("", text).strip()


# ── chapter grouping ──────────────────────────────────────────────────────────

def chapters(rows):
    """Group a verse stream into (book, chapter, [(n, text)…]) units.

    Grouping is on change of (book, chapter) in STREAM order, never on sorting
    a collected list — in a chronological edition a book legitimately recurs
    (Genesis, then Job, then Genesis again), and each recurrence must stay
    where the sequence put it.
    """
    cur_key, buf = None, []
    for book, chap, num, text in rows:
        key = (book, chap)
        if key != cur_key:
            if cur_key:
                yield cur_key[0], cur_key[1], buf
            cur_key, buf = key, []
        buf.append((num, text))
    if cur_key:
        yield cur_key[0], cur_key[1], buf


# ── writers ───────────────────────────────────────────────────────────────────

class Writer:
    """A writer turns chapter units into one file in the shop's design.

    Subclasses receive `styles`: the role -> shop's-own-name map loaded from
    --styles, plus whatever extra keys that file carries.
    """

    def __init__(self, styles, template=None, meta=None):
        self.styles = styles
        self.template = template
        self.meta = meta or {}

    def render(self, units, out_path):
        raise NotImplementedError


class TaggedTextWriter(Writer):
    """InDesign Tagged Text.

    Placed into the shop's document, each <ParaStyle:…> resolves against their
    style definitions, so their fonts, spacing, chapter rules and ornaments
    apply without us knowing anything about them. InDesign requires UTF-16LE
    for the <UNICODE-WIN> header; writing UTF-8 produces mojibake with no error
    message, which is why the encoding is pinned here and not left to the OS.
    """

    HEADER = "<UNICODE-WIN>\n<vsn:8><fset:InDesign-Roman>\n"

    @staticmethod
    def esc(s):
        # Backslash first, or the escapes introduced below get double-escaped.
        return s.replace("\\", "\\\\").replace("<", "\\<").replace(">", "\\>")

    def para(self, role, text):
        style = self.styles.get(role)
        if not style:
            return ""
        return f"<ParaStyle:{self.esc(style)}>{text}\n"

    def render(self, units, out_path):
        vn_style = self.styles.get("verse_number")
        parts = [self.HEADER]
        last_book = None
        first = True
        for book, chap, verses in units:
            if book != last_book:
                parts.append(self.para("book_title", self.esc(book)))
                last_book = book
            elif not first:
                parts.append(self.para("separator",
                                       self.esc(self.styles.get("separator_text", ""))))
            first = False
            parts.append(self.para("chapter_head", self.esc(f"{book} {chap}")))
            body = []
            for n, t in verses:
                if vn_style:
                    body.append(f"<CharStyle:{self.esc(vn_style)}>{self.esc(n)}"
                                f"<CharStyle:>{self.esc(t)}")
                else:
                    body.append(self.esc(t))
            parts.append(self.para("body", " ".join(body)))
        with open(out_path, "w", encoding="utf-16-le") as fh:
            fh.write("﻿" + "".join(parts))
        return out_path


class HtmlWriter(Writer):
    """The shop's HTML shell, filled fragment by fragment.

    Their file supplies @font-face, page CSS and ornament markup. We only
    substitute tokens inside the fragments they marked, so a separator they
    drew in SVG comes back out byte-identical.
    """

    FRAG = re.compile(r"<!--\s*CR:(\w+)\s*-->(.*?)<!--\s*/CR:\1\s*-->", re.S)

    def fragments(self):
        found = {m.group(1).lower(): m.group(2)
                 for m in self.FRAG.finditer(self.template)}
        missing = [k for k in ("chapter", "verse") if k not in found]
        if missing:
            raise SystemExit(
                "Template is missing required fragment(s): "
                + ", ".join(f"<!--CR:{m.upper()}--> … <!--/CR:{m.upper()}-->"
                            for m in missing)
                + "\nSee docs/printer-templates.md for a worked example.")
        return found

    def render(self, units, out_path):
        frag = self.fragments()
        out = []
        last_book = None
        first = True
        for book, chap, verses in units:
            if book != last_book and "book" in frag:
                out.append(frag["book"].replace("{{BOOK_NAME}}",
                                                htmlmod.escape(book)))
            elif not first and "separator" in frag:
                out.append(frag["separator"])
            if book != last_book:
                last_book = book
            first = False
            vs = "".join(
                frag["verse"].replace("{{N}}", htmlmod.escape(n))
                             .replace("{{TEXT}}", htmlmod.escape(t))
                for n, t in verses)
            out.append(frag["chapter"]
                       .replace("{{CHAPTER_NAME}}", htmlmod.escape(f"{book} {chap}"))
                       .replace("{{BOOK_NAME}}", htmlmod.escape(book))
                       .replace("{{CHAPTER_NUMBER}}", htmlmod.escape(chap))
                       .replace("{{VERSES}}", vs))
        page = self.FRAG.sub("", self.template)      # fragments are prototypes
        page = page.replace("{{BODY}}", "".join(out))
        page = page.replace("{{TITLE}}", htmlmod.escape(self.meta.get("title", "")))
        with open(out_path, "w", encoding="utf-8") as fh:
            fh.write(page)
        return out_path


class DocxWriter(Writer):
    """Their .docx opened as a template; paragraphs appended with their styles.

    python-docx raises KeyError for a style name the template does not define,
    which is the right behaviour — a silent fallback to Normal would return a
    document that looks like ours, not theirs.
    """

    def render(self, units, out_path):
        try:
            from docx import Document
        except ImportError:
            raise SystemExit("The docx format needs python-docx: "
                             "pip3 install python-docx --break-system-packages")
        doc = Document(self.template) if self.template else Document()
        last_book = None
        first = True
        for book, chap, verses in units:
            if book != last_book:
                self._para(doc, "book_title", book)
                last_book = book
            elif not first:
                self._para(doc, "separator", self.styles.get("separator_text", ""))
            first = False
            self._para(doc, "chapter_head", f"{book} {chap}")
            p = doc.add_paragraph(style=self.styles.get("body") or None)
            for n, t in verses:
                if self.styles.get("verse_number"):
                    run = p.add_run(n)
                    run.font.superscript = True
                p.add_run(t + " ")
        doc.save(out_path)
        return out_path

    def _para(self, doc, role, text):
        style = self.styles.get(role)
        if style and text:
            doc.add_paragraph(text, style=style)


class IdmlWriter(Writer):
    """Deliberately not implemented.

    Generating an IDML means we author the InDesign document: master pages,
    text frames, style definitions, margins. That is the opposite of filling
    the shop's template — it hands us their design decisions, which is both
    more work and worse. If a shop insists on IDML, the honest answer is that
    they place our Tagged Text into their own IDML.
    """

    def render(self, units, out_path):
        raise SystemExit(
            "IDML is not supported by design: generating one would mean us "
            "authoring the InDesign document rather than filling theirs. Ask "
            "the shop to place the --format tagged output into their template; "
            "it carries their styles through unchanged.")


WRITERS = {"tagged": TaggedTextWriter, "html": HtmlWriter,
           "docx": DocxWriter, "idml": IdmlWriter}


# ── style map ─────────────────────────────────────────────────────────────────

def load_styles(path):
    """role -> the shop's own style name. YAML if PyYAML is present, else JSON.

    Deliberately forgiving about the file format and strict about the role
    names: a typo in a role is silently dropped work, so it is an error.
    """
    if not path:
        return {}
    raw = open(path, encoding="utf-8").read()
    data = None
    if path.endswith((".yml", ".yaml")):
        try:
            import yaml
            data = yaml.safe_load(raw)
        except ImportError:
            raise SystemExit("A .yml style map needs PyYAML, or write it as "
                             ".json: pip3 install pyyaml --break-system-packages")
    else:
        data = json.loads(raw)
    styles = data.get("styles", data) or {}
    unknown = [k for k in styles
               if k not in ROLES and k not in ("separator_text",)]
    if unknown:
        raise SystemExit(
            f"Unknown role(s) in {path}: {', '.join(sorted(unknown))}. "
            f"Known roles: {', '.join(ROLES)} (plus separator_text).")
    return styles


# ── self-test ─────────────────────────────────────────────────────────────────

FIXTURE = [
    ("Deuteronomy", "34", "12", "And in all that mighty hand."),
    ("Psalm", "90", "1", "Lord, thou hast been our dwelling place."),
    ("Psalm", "90", "2", "Before the mountains were brought forth."),
    ("Psalm", "91", "1", "He that dwelleth in the secret place."),
    ("Joshua", "1", "1", "Now after the death of Moses <the servant>."),
]


def self_test():
    """Prove the writers without BaseX, so a template can be checked in
    isolation from a corpus problem — the two failure modes look identical
    from the outside otherwise."""
    import tempfile
    units = list(chapters(iter(FIXTURE)))
    assert [(b, c) for b, c, _ in units] == [
        ("Deuteronomy", "34"), ("Psalm", "90"), ("Psalm", "91"),
        ("Joshua", "1")], units
    assert len(units[1][2]) == 2, "chapter grouping lost a verse"

    styles = {"book_title": "Book Title", "chapter_head": "Chapter Head",
              "body": "Body Text", "verse_number": "Verse Num",
              "separator": "Ornament", "separator_text": "❦"}
    d = tempfile.mkdtemp()

    p = TaggedTextWriter(styles).render(units, os.path.join(d, "t.txt"))
    tagged = open(p, encoding="utf-16-le").read().lstrip("﻿")
    assert tagged.startswith("<UNICODE-WIN>"), "InDesign header missing"
    assert "<ParaStyle:Chapter Head>Psalm 90" in tagged
    assert "<CharStyle:Verse Num>1<CharStyle:>" in tagged
    assert "\\<the servant\\>" in tagged, "angle brackets not escaped for InDesign"
    assert "<ParaStyle:Ornament>❦" in tagged, "chapter separator not emitted"

    tpl = ("<html><head><style>@font-face{}</style></head><body>{{BODY}}"
           "<!--CR:BOOK--><h1>{{BOOK_NAME}}</h1><!--/CR:BOOK-->"
           "<!--CR:SEPARATOR--><p class=orn>&#10086;</p><!--/CR:SEPARATOR-->"
           "<!--CR:CHAPTER--><h2>{{CHAPTER_NAME}}</h2><p>{{VERSES}}</p><!--/CR:CHAPTER-->"
           "<!--CR:VERSE--><sup>{{N}}</sup>{{TEXT}} <!--/CR:VERSE-->"
           "</body></html>")
    p = HtmlWriter({}, template=tpl).render(units, os.path.join(d, "t.html"))
    got = open(p, encoding="utf-8").read()
    assert "<h1>Deuteronomy</h1>" in got
    assert "<h2>Psalm 90</h2>" in got
    assert "<sup>2</sup>Before the mountains" in got
    assert "&#10086;" in got, "separator fragment not reproduced verbatim"
    assert "CR:CHAPTER" not in got, "fragment prototypes left in the output"
    assert "&lt;the servant&gt;" in got, "verse text not escaped"

    assert clean("In the beginning. 1.3 substance: or, cattle") == "In the beginning."
    print("self-test OK — writers, grouping, escaping and footnote strip")
    return 0


# ── main ──────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--translation", help="BaseX document id, e.g. bible-kjv-1611")
    ap.add_argument("--ordering", default="canonical", choices=sorted(ORDERINGS))
    ap.add_argument("--format", default="tagged", choices=sorted(WRITERS))
    ap.add_argument("--template", help="the shop's HTML shell or .docx")
    ap.add_argument("--styles", help="role -> the shop's style names (.yml/.json)")
    ap.add_argument("--out", help="output file")
    ap.add_argument("--books", help="comma-separated bookName filter (taster only)")
    ap.add_argument("--limit", type=int, help="stop after N verses (dry runs)")
    ap.add_argument("--title", default="")
    ap.add_argument("--self-test", action="store_true",
                    help="check the writers without touching BaseX")
    a = ap.parse_args()

    if a.self_test:
        return self_test()
    if not (a.translation and a.out):
        ap.error("--translation and --out are required (or use --self-test)")
    if a.format in ("html", "docx") and not a.template:
        ap.error(f"--format {a.format} needs the shop's --template file")
    if a.format in ("tagged", "docx") and not a.styles:
        ap.error(f"--format {a.format} needs --styles naming their paragraph styles")

    styles = load_styles(a.styles)
    template = None
    if a.template and a.format == "html":
        template = open(a.template, encoding="utf-8").read()
    elif a.template:
        template = a.template            # docx: a path, opened by python-docx

    books = [b.strip() for b in a.books.split(",")] if a.books else None
    try:
        rows = stream_verses(a.translation, a.ordering, books=books, limit=a.limit)
        units = list(chapters(rows))
    except CorpusError as e:
        sys.exit(str(e))

    if not units:
        sys.exit(f"No verses matched. Check --books names (they are @bookName "
                 f"values such as 'Psalm', not codes such as 'PSA').")

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    writer = WRITERS[a.format](styles, template=template,
                               meta={"title": a.title or a.translation})
    path = writer.render(units, a.out)
    verses_out = sum(len(v) for _, _, v in units)
    print(f"{path}: {len(units)} chapters, {verses_out} verses, "
          f"{a.ordering} order, {a.format} format")
    return 0


if __name__ == "__main__":
    sys.exit(main())
